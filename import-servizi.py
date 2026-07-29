#!/usr/bin/env python3
"""
Importa i LAVORI (servizi) dal file mensile "GIUGNO 2026_.xlsx" e genera
js/servizi-reali.js, sostituendo integralmente i servizi caricati in precedenza.

Ogni riga dell'Excel = una persona su un lavoro (data, cliente, indirizzo,
attività, orario, ore totali, nominativo, società). Le righe identiche per
data + cliente + indirizzo + attività + orario vengono raggruppate in un unico
servizio con N risorse e l'elenco delle persone effettivamente assegnate.

I nominativi vengono agganciati all'anagrafica (js/dipendenti-reali.js):
chi non c'è viene marcato come risorsa esterna (interinali, ditte esterne).

Uso:  pip3 install openpyxl ; python3 import-servizi.py
"""
import json
import os
import re
import unicodedata
import datetime
from collections import defaultdict

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
EXCEL = os.path.join(HERE, "GIUGNO 2026_.xlsx")
DIPENDENTI = os.path.join(HERE, "js", "dipendenti-reali.js")
OUT = os.path.join(HERE, "js", "servizi-reali.js")

# Indici di colonna nel foglio (0-based), come da intestazione del file.
COL = {
    "data": 0, "cliente": 1, "indirizzo": 2, "attivita": 3,
    "ora_in1": 4, "ora_fi1": 5, "ora_in2": 6, "ora_fi2": 7,
    "tot_ore": 8, "nominativo": 9, "societa": 10,
    "mezzi": 11, "extra": 12, "ore_viaggio": 13, "rimborsi": 14, "note": 15,
}

# Attività: normalizzazione delle diciture dell'Excel (maiuscole, plurali, refusi).
ATTIVITA_MAP = {
    "PULIZIA": "Pulizie", "PULIZIE": "Pulizie", "PULIZIA VETRI": "Pulizie vetri",
    "PULIZIA/PRESIDIO": "Pulizie e presidio", "PULIZIE/PRESISIO": "Pulizie e presidio",
    "PULIZIE/PRESIDIO": "Pulizie e presidio", "PRESIDIO": "Presidio",
    "PULIZIA TOMBA": "Pulizie",
    "FACCHINAGGIO": "Facchinaggio", "FACCHINAGGO": "Facchinaggio",
    "FACCH/PULIZIE": "Facchinaggio e pulizie", "IMBIANCATURA/PULIZIE": "Imbiancatura e pulizie",
    "FACCH/SUPPORTO MONT": "Facchinaggio e montaggio",
    "FACCHINAGGIO/SUPPORTO": "Facchinaggio e montaggio",
    "SUPP. MONTAGGIO": "Supporto montaggio", "SUPP.MONTAGGIO": "Supporto montaggio",
    "MONTAGGIO": "Montaggio", "SMONTAGGIO": "Smontaggio",
    "IMBIANCATURA": "Imbiancatura", "A.ASCENSORI": "Assist. ascensori",
    "ASSISTENZA ASCENSORI": "Assist. ascensori",
    "MAGAZZINIERE": "Magazziniere", "MAGAZZINO": "Magazziniere",
    "CONFEZIONAMENTO": "Confezionamento",
    "CONSEGNA MATERIALE IMBALLAGGIO": "Consegna materiale",
    "TRASPORTO": "Trasporto", "TRASLOCO": "Trasloco", "TRASLOCHI": "Trasloco",
    "CORSO IN AZIENDA": "Corso in azienda", "ASSISTENZA": "Assistenza",
}

# Soglie (giorni distinti nel mese, per cantiere = cliente+indirizzo) -> categoria.
# Sostituiscono i colori del vecchio file: la ricorrenza è ricavata dai dati.
SOGLIE_CAT = [(15, "verde"), (6, "azzurro"), (2, "viola"), (0, "bianco")]


# ---------------------------------------------------------------- utilità ---
def testo(v) -> str:
    return "" if v is None else str(v).strip()


def compatta(v: str) -> str:
    """Spazi multipli ridotti a uno (i dati hanno molti doppi spazi finali)."""
    return re.sub(r"\s+", " ", testo(v)).strip()


def data_iso(v) -> str:
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    s = testo(v)
    m = re.match(r"^(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})$", s)
    if m:
        g, me, a = (int(x) for x in m.groups())
        a = a + 2000 if a < 100 else a
        try:
            return datetime.date(a, me, g).isoformat()
        except ValueError:
            return ""
    return ""


def ora_hhmm(v) -> str:
    """Orario "HH:MM". Gestisce time, datetime (date fittizia 1900) e testo sporco."""
    if isinstance(v, datetime.time):
        return f"{v.hour:02d}:{v.minute:02d}"
    if isinstance(v, datetime.datetime):
        return f"{v.hour:02d}:{v.minute:02d}"
    s = testo(v)
    m = re.match(r"^(\d{1,2})[:.](\d{1,2})", s)
    if m:
        h, mi = int(m.group(1)), int(m.group(2))
        if 0 <= h <= 23 and 0 <= mi <= 59:
            return f"{h:02d}:{mi:02d}"
    m = re.match(r"^(\d{1,2})::?$", s)  # refusi tipo "17::"
    if m and int(m.group(1)) <= 23:
        return f"{int(m.group(1)):02d}:00"
    return ""


def ore_decimali(v):
    """Colonna TOTALE ORE -> ore in decimale (8:30 = 8.5). None se non leggibile."""
    if isinstance(v, datetime.time):
        return v.hour + v.minute / 60
    if isinstance(v, datetime.datetime):
        return v.hour + v.minute / 60
    if isinstance(v, (int, float)):
        return float(v) if 0 < float(v) <= 24 else None
    s = testo(v).replace(",", ".")
    m = re.match(r"^(\d{1,2})[:.](\d{1,2})$", s)
    if m:
        return int(m.group(1)) + int(m.group(2)) / 60
    try:
        n = float(s)
        return n if 0 < n <= 24 else None
    except ValueError:
        return None


def numero(v):
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, datetime.time):
        return v.hour + v.minute / 60
    s = testo(v).replace(",", ".").replace("€", "").strip()
    try:
        return float(s)
    except ValueError:
        return 0.0


def minuti(hhmm: str):
    if not hhmm:
        return None
    h, m = hhmm.split(":")
    return int(h) * 60 + int(m)


def durata_ore(inizio: str, fine: str, pausa_da: str, pausa_a: str):
    """Ore fra inizio e fine (gestisce il turno che passa la mezzanotte), meno la pausa."""
    a, b = minuti(inizio), minuti(fine)
    if a is None or b is None:
        return None
    tot = (b - a) % (24 * 60)
    if tot == 0:
        return None
    pa, pb = minuti(pausa_da), minuti(pausa_a)
    if pa is not None and pb is not None:
        pausa = (pb - pa) % (24 * 60)
        if 0 < pausa < tot:
            tot -= pausa
    return tot / 60


def token_nome(s: str):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return set(re.sub(r"[^A-Za-z ]", " ", s).upper().split())


def titolo_nome(s: str) -> str:
    return " ".join(p.capitalize() for p in compatta(s).split())


def cat_da_giorni(n: int) -> str:
    for soglia, cat in SOGLIE_CAT:
        if n >= soglia:
            return cat
    return "bianco"


# --------------------------------------------------------------- anagrafica ---
def carica_dipendenti():
    """[(nome, id, set_token)] dall'anagrafica generata da import-dipendenti.py."""
    if not os.path.exists(DIPENDENTI):
        print("⚠️  js/dipendenti-reali.js non trovato: i nominativi resteranno non agganciati.")
        return []
    src = open(DIPENDENTI, encoding="utf-8").read()
    if "window.GL_SEED = " not in src:
        return []
    grezzo = src.split("window.GL_SEED = ", 1)[1].rsplit(";", 1)[0].strip()
    return [(d["nome"], d["id"], token_nome(d["nome"])) for d in json.loads(grezzo)]


def aggancia(nome: str, dipendenti):
    """Nominativo Excel -> (nome da mostrare, id dipendente | None).

    Prima il confronto esatto sui token (l'ordine nome/cognome cambia fra i file),
    poi una similarità alta (>= 0.5 di Jaccard con almeno 2 token) per i refusi.
    """
    t = token_nome(nome)
    if not t:
        return titolo_nome(nome), None
    for n, i, td in dipendenti:
        if t == td:
            return n, i
    migliore, punteggio, comuni = None, 0.0, 0
    for n, i, td in dipendenti:
        inter = len(t & td)
        j = inter / len(t | td) if (t | td) else 0
        if j > punteggio:
            migliore, punteggio, comuni = (n, i), j, inter
    if migliore and punteggio >= 0.5 and comuni >= 2:
        return migliore
    return titolo_nome(nome), None


# --------------------------------------------------------------------- main ---
def leggi_righe():
    wb = openpyxl.load_workbook(EXCEL, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    righe, scartate = [], 0
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue  # intestazione
        if not r or all(v is None or testo(v) == "" for v in r):
            continue
        val = lambda k: r[COL[k]] if COL[k] < len(r) else None  # noqa: E731
        data = data_iso(val("data"))
        nome = compatta(val("nominativo"))
        if not data or not nome or nome == "/":
            scartate += 1
            continue
        inizio = ora_hhmm(val("ora_in1")) or ora_hhmm(val("ora_in2"))
        fine = ora_hhmm(val("ora_fi2")) or ora_hhmm(val("ora_fi1"))
        pausa_da, pausa_a = ora_hhmm(val("ora_fi1")), ora_hhmm(val("ora_in2"))
        pausa = f"{pausa_da}–{pausa_a}" if (pausa_da and pausa_a and fine != pausa_da) else ""
        ore = ore_decimali(val("tot_ore"))
        stimate = False
        if ore is None:
            ore = durata_ore(inizio, fine, pausa_da, pausa_a)
            stimate = ore is not None
        attivita_raw = compatta(val("attivita")).upper()
        righe.append({
            "data": data,
            "cliente": compatta(val("cliente")) or "Cliente da definire",
            "indirizzo": compatta(val("indirizzo")),
            "attivita": ATTIVITA_MAP.get(attivita_raw, titolo_nome(attivita_raw) or "Servizio"),
            "ora": inizio,
            "oraFine": fine,
            "pausa": pausa,
            "ore": round(ore, 2) if ore else None,
            "oreStimate": stimate,
            "nominativo": nome,
            "societa": compatta(val("societa")).upper(),
            "oreViaggio": numero(val("ore_viaggio")),
            "rimborsi": numero(val("rimborsi")),
            "note": compatta(val("note")),
        })
    wb.close()
    return righe, scartate


def main():
    dipendenti = carica_dipendenti()
    righe, scartate = leggi_righe()
    print(f"Righe lette: {len(righe)} (scartate senza data/nominativo: {scartate})")

    # Categoria per cantiere: quanti giorni distinti del mese è presente.
    giorni_cantiere = defaultdict(set)
    for r in righe:
        giorni_cantiere[(r["cliente"].upper(), r["indirizzo"].upper())].add(r["data"])

    # Raggruppa le righe identiche: stesso giorno, cantiere, attività e orario.
    gruppi = {}
    for r in righe:
        chiave = (r["data"], r["cliente"].upper(), r["indirizzo"].upper(),
                  r["attivita"], r["ora"], r["oraFine"])
        gruppi.setdefault(chiave, []).append(r)

    cache_nomi = {}
    servizi, esterni, senza_ore = [], set(), 0
    for n, (chiave, lista) in enumerate(sorted(gruppi.items()), start=1):
        primo = lista[0]
        persone = []
        for r in lista:
            if r["nominativo"] not in cache_nomi:
                cache_nomi[r["nominativo"]] = aggancia(r["nominativo"], dipendenti)
            nome, dip_id = cache_nomi[r["nominativo"]]
            if not dip_id:
                esterni.add(nome)
            persone.append({"n": nome, "id": dip_id})
        persone.sort(key=lambda p: p["n"])

        ore = next((r["ore"] for r in lista if r["ore"]), None)
        if ore is None:
            senza_ore += 1
        note = next((r["note"] for r in lista if r["note"]), "")
        societa = next((r["societa"] for r in lista if r["societa"]), "")
        giorni = len(giorni_cantiere[(primo["cliente"].upper(), primo["indirizzo"].upper())])

        servizi.append({
            "id": f"s{n}",
            "data": primo["data"],
            "cliente": primo["cliente"],
            "indirizzo": primo["indirizzo"],
            "attivita": primo["attivita"],
            "ora": primo["ora"],
            "oraFine": primo["oraFine"],
            "pausa": primo["pausa"],
            "ore": ore or 0,
            "oreStimate": any(r["oreStimate"] for r in lista) or ore is None,
            "risorse": len(persone),
            "persone": persone,
            "cat": cat_da_giorni(giorni),
            "societa": societa,
            "oreViaggio": round(sum(r["oreViaggio"] for r in lista), 2),
            "rimborsi": round(sum(r["rimborsi"] for r in lista), 2),
            "note": note,
        })

    servizi.sort(key=lambda s: (s["data"], s["cliente"].upper(), s["ora"] or "99:99"))
    for i, s in enumerate(servizi, start=1):
        s["id"] = f"s{i}"

    mesi = sorted({s["data"][:7] for s in servizi})
    ore_tot = sum(s["ore"] * s["risorse"] for s in servizi)
    meta = {
        "mese": mesi[0] if len(mesi) == 1 else ", ".join(mesi),
        "righe": len(righe),
        "servizi": len(servizi),
        "risorse": sum(s["risorse"] for s in servizi),
        "oreTotali": round(ore_tot, 2),
        "esterni": sorted(esterni),
    }

    with open(OUT, "w", encoding="utf-8") as f:
        f.write("/* Servizi REALI importati da 'GIUGNO 2026_.xlsx' — generato da import-servizi.py.\n")
        f.write("   Una riga Excel = una persona su un lavoro; le righe identiche (giorno, cliente,\n")
        f.write("   indirizzo, attività, orario) sono raggruppate in un servizio con N risorse.\n")
        f.write("   Categoria per ricorrenza del cantiere nel mese: verde >= 15 giorni, azzurro >= 6,\n")
        f.write("   viola >= 2, bianco = spot.\n")
        f.write("   ATTENZIONE: contiene nomi di persone, clienti e indirizzi. NON pubblicare in chiaro. */\n")
        f.write("window.GL_SERVIZI_META = " + json.dumps(meta, ensure_ascii=False) + ";\n")
        f.write("window.GL_SERVIZI_REALI = [\n")
        f.write(",\n".join(" " + json.dumps(s, ensure_ascii=False) for s in servizi))
        f.write("\n];\n")

    agganciati = sum(1 for v in cache_nomi.values() if v[1])
    print(f"Generato {OUT}")
    print(f"Servizi: {len(servizi)} — risorse impiegate: {meta['risorse']} — ore totali: {meta['oreTotali']}")
    print(f"Nominativi distinti: {len(cache_nomi)} — agganciati all'anagrafica: {agganciati} — esterni: {len(esterni)}")
    if senza_ore:
        print(f"⚠️  Servizi senza ore leggibili: {senza_ore} (contano 0 ore nei calcoli)")
    if esterni:
        print("   Esterni/non in anagrafica: " + ", ".join(sorted(esterni)))


if __name__ == "__main__":
    main()
